"""Bounded, inert native Office text extraction; run parsing in an isolated process."""

import io
import json
import math
import os
import re
import resource
import struct
import sys
import zipfile
from datetime import date, datetime
from pathlib import PurePosixPath

MAX_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 40 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 2000
MAX_TEXT = 120000
MAX_UNIT_TEXT = 16000
MAX_UNITS = 25
MAX_ROWS = 2000
MAX_CELLS = 10000
MIME_BY_FORMAT = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "xls": "application/vnd.ms-excel",
}
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
S_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
HIDDEN_CONTENT_ERROR = (
    'The workbook contains populated hidden sheets, rows or columns. Review and unhide the intended material, '
    'or save a clean copy containing only the intended visible purchase order before uploading.'
)


class OfficeReadError(ValueError):
    """The complete document cannot be represented safely within intake bounds."""


def _archive(content):
    # Bound central-directory allocation before ZipFile materializes its entries
    # in the API process. Document parsing itself runs in the resource-limited child.
    end = content.rfind(b"PK\x05\x06", max(0, len(content) - 65557))
    if end < 0 or end + 22 > len(content):
        raise OfficeReadError("The Office archive is damaged.")
    (
        _,
        disk,
        directory_disk,
        disk_entries,
        entries,
        directory_size,
        directory_offset,
        _,
    ) = struct.unpack_from("<4s4H2IH", content, end)
    if (
        disk
        or directory_disk
        or disk_entries != entries
        or entries > MAX_ARCHIVE_MEMBERS
        or directory_size > 2 * 1024 * 1024
        or directory_offset + directory_size > end
    ):
        raise OfficeReadError("The Office archive exceeds safe directory limits.")
    archive = zipfile.ZipFile(io.BytesIO(content))
    members = archive.infolist()
    names = [member.filename for member in members]
    if len(members) > MAX_ARCHIVE_MEMBERS or len(names) != len(set(names)):
        raise OfficeReadError("The Office archive contains too many or duplicate entries.")
    total = 0
    for member in members:
        name = member.filename.replace("\\", "/")
        if name.startswith("/") or ".." in PurePosixPath(name).parts or member.flag_bits & 1:
            raise OfficeReadError("Encrypted or unsafe Office archives are not supported.")
        total += member.file_size
        if total > MAX_EXPANDED_BYTES or member.file_size > 16 * 1024 * 1024:
            raise OfficeReadError("The expanded Office document is too large. Split it before uploading.")
        if member.file_size > 1024 * 1024 and member.file_size > max(member.compress_size, 1) * 200:
            raise OfficeReadError("The Office archive exceeds safe decompression limits.")
        lower = name.casefold()
        if any(value in lower for value in ("vbaproject", "/macrosheets/", "/embeddings/", "/activex/")):
            raise OfficeReadError("Macros and embedded executable objects are not supported.")
    return archive


def detect_intake_format(content, filename=None):
    """Check actual container signatures and required parts, never trust supplied MIME."""
    if not 0 < len(content) <= MAX_BYTES:
        raise OfficeReadError("Each document must be no larger than 10 MB.")
    if content.startswith(b"%PDF-"):
        kind = "pdf"
    elif content.startswith(OLE_MAGIC):
        kind = "xls"
    elif content.startswith(b"PK\x03\x04"):
        try:
            with _archive(content) as archive:
                names = set(archive.namelist())
                if "word/document.xml" in names and "xl/workbook.xml" not in names:
                    kind = "docx"
                elif "xl/workbook.xml" in names and "word/document.xml" not in names:
                    kind = "xlsx"
                else:
                    raise OfficeReadError("This archive is not a supported Word or Excel document.")
                if "[Content_Types].xml" not in names:
                    raise OfficeReadError("The Office document is missing required content metadata.")
        except (zipfile.BadZipFile, KeyError, OSError) as exc:
            raise OfficeReadError("The Office document is damaged or unsupported.") from exc
    else:
        raise OfficeReadError("Upload a PDF, DOCX, XLSX or XLS document.")
    if filename and os.path.splitext(filename)[1].casefold() != "." + kind:
        raise OfficeReadError("The filename extension does not match the document contents.")
    return kind


def _xml(archive, name):
    from lxml import etree

    content = archive.read(name)
    declaration_scan = content.replace(b'\x00', b'').upper()
    if b"<!DOCTYPE" in declaration_scan or b"<!ENTITY" in declaration_scan:
        raise OfficeReadError("Office XML entity declarations are not supported.")
    root = etree.fromstring(content, parser=etree.XMLParser(resolve_entities=False, no_network=True))
    for node in root.iter():
        if not isinstance(node.tag, str) or not node.tag.startswith(f'{{{W_NS}}}'):
            continue
        tag = node.tag.rsplit('}', 1)[-1]
        if (
            tag in {'ins', 'del', 'cellDel', 'cellIns', 'cellMerge', 'numberingChange'}
            or tag.endswith('PrChange')
            or tag.startswith(('moveFrom', 'moveTo', 'customXmlDel', 'customXmlIns', 'customXmlMove'))
        ):
            raise OfficeReadError(
                'Word tracked changes are not supported. Accept or reject the revisions and save a clean DOCX, '
                'or export the final document to PDF before uploading.'
            )
    return root


class _TextUnits:
    def __init__(self):
        self.units = []
        self.labels = []
        self.total = 0
        self.rows = 0
        self.group = None

    def add(self, locator, value, *, group="Document"):
        value = str(value).strip()
        if not value:
            return
        text = f"[{locator}] {value}"
        self.total += len(text) + 1
        self.rows += 1
        if len(text) > MAX_UNIT_TEXT or self.total > MAX_TEXT or self.rows > MAX_ROWS:
            raise OfficeReadError("The document contains too much text or too many rows. Split it before uploading.")
        if not self.units or self.group != group or len(self.units[-1]) + len(text) + 1 > MAX_UNIT_TEXT:
            if len(self.units) >= MAX_UNITS:
                raise OfficeReadError("The document exceeds 25 evidence sections. Split it before uploading.")
            self.units.append("")
            self.labels.append(f"{group} — section {len(self.units)}")
            self.group = group
        self.units[-1] += ("\n" if self.units[-1] else "") + text

    def result(self, kind, warnings):
        if not self.units:
            raise OfficeReadError("No readable text was found. Export image-only documents to PDF and upload the PDF.")
        return {
            "format": kind,
            "units": self.units,
            "labels": self.labels,
            "warnings": warnings[:6],
        }


def _read_docx(archive):
    text = _TextUnits()
    warnings = []
    names = archive.namelist()
    # These parts can contain material missing from simple paragraph extraction;
    # decline rather than silently treating an incomplete source as complete.
    if any("/afchunk" in name.casefold() for name in names):
        raise OfficeReadError("Word imported content is unsupported. Export this document to PDF first.")
    if any(name.startswith("word/media/") for name in names):
        warnings.append(
            "Embedded images are not read in Word uploads. Check them in the original or upload a PDF for visual extraction."
        )
    if any(
        name.startswith("word/") and name.endswith(".rels") and b'TargetMode="External"' in archive.read(name)
        for name in names
    ):
        warnings.append("External links were not opened; only text stored in the document was read.")
    counters = {"p": 0, "tbl": 0}

    def paragraph(node):
        fragments = []
        for child in node.iter():
            if any(parent.tag == f"{{{W_NS}}}del" for parent in child.iterancestors()):
                continue
            if child.tag == f"{{{W_NS}}}t":
                fragments.append(child.text or "")
            elif child.tag in (f"{{{W_NS}}}tab", f"{{{W_NS}}}br", f"{{{W_NS}}}cr"):
                fragments.append(" ")
        return "".join(fragments)

    def walk(parent, label):
        for node in parent:
            tag = node.tag.rsplit("}", 1)[-1] if isinstance(node.tag, str) else ""
            if tag == "p":
                counters["p"] += 1
                text.add(f'{label} paragraph {counters["p"]}', paragraph(node), group=label)
            elif tag == "tbl":
                counters["tbl"] += 1
                number = counters["tbl"]
                for index, row in enumerate(node.findall(f"{{{W_NS}}}tr"), 1):
                    cells = []
                    for cell in row.findall(f"{{{W_NS}}}tc"):
                        cells.append(" / ".join(paragraph(p) for p in cell.iter(f"{{{W_NS}}}p")))
                    text.add(
                        f"{label} table {number} row {index}",
                        " | ".join(cells),
                        group=label,
                    )
            elif tag == "altChunk":
                raise OfficeReadError("Word imported content is unsupported. Export this document to PDF first.")
            elif tag not in ("sectPr", "del"):
                walk(node, label)

    for name in sorted(name for name in names if re.fullmatch(r"word/header\d+\.xml", name)):
        walk(_xml(archive, name), os.path.basename(name).replace(".xml", "").title())
    root = _xml(archive, "word/document.xml")
    body = root.find(f"{{{W_NS}}}body")
    if body is None:
        raise OfficeReadError("The Word document has no readable body.")
    walk(body, "Document")
    for name in ("word/footnotes.xml", "word/endnotes.xml"):
        if name in names:
            walk(_xml(archive, name), os.path.basename(name).replace(".xml", "").title())
    for name in sorted(name for name in names if re.fullmatch(r"word/footer\d+\.xml", name)):
        walk(_xml(archive, name), os.path.basename(name).replace(".xml", "").title())
    return text.result("docx", warnings)


def _display(value, number_format=None):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if not math.isfinite(value):
            raise OfficeReadError("The workbook contains a nonfinite numeric value.")
        if value.is_integer():
            value = int(value)
    # Keep simple printed identifier padding (000123) without turning locale,
    # currency or accounting formats into guessed units or prices.
    if isinstance(value, int) and number_format and re.fullmatch(r"0{2,}", number_format):
        return str(value).zfill(len(number_format))
    return str(value)


def _validate_xlsx_visibility(archive):
    from openpyxl.utils.cell import column_index_from_string

    for name in archive.namelist():
        if not re.fullmatch(r'xl/worksheets/[^/]+\.xml', name):
            continue
        root = _xml(archive, name)
        hidden_columns = []
        for column in root.findall(f'{{{S_NS}}}cols/{{{S_NS}}}col'):
            if column.get('hidden', '').lower() in ('1', 'true') or float(column.get('width', '1')) <= 0:
                hidden_columns.append((int(column.get('min')), int(column.get('max'))))
        hidden_rows = set()
        for row in root.findall(f'{{{S_NS}}}sheetData/{{{S_NS}}}row'):
            if row.get('hidden', '').lower() in ('1', 'true') or float(row.get('ht', '1')) <= 0:
                hidden_rows.add(int(row.get('r')))
        dimensions = root.find(f'{{{S_NS}}}sheetFormatPr')
        default_hidden = dimensions is not None and dimensions.get('zeroHeight', '').lower() in ('1', 'true')
        # Inspect actual populated cells even when workbook dimensions are false.
        for cell in root.findall(f'{{{S_NS}}}sheetData/{{{S_NS}}}row/{{{S_NS}}}c'):
            has_value = any(
                node.tag == f'{{{S_NS}}}f' or node.tag in (f'{{{S_NS}}}v', f'{{{S_NS}}}t') and node.text
                for node in cell.iter()
            )
            if not has_value:
                continue
            coordinate = re.fullmatch(r'([A-Z]+)([0-9]+)', cell.get('r', ''))
            if coordinate is None:
                raise OfficeReadError('The workbook contains invalid cell coordinates.')
            column, row = column_index_from_string(coordinate.group(1)), int(coordinate.group(2))
            if default_hidden or row in hidden_rows or any(start <= column <= end for start, end in hidden_columns):
                raise OfficeReadError(HIDDEN_CONTENT_ERROR)


def _read_xlsx(content, archive):
    from openpyxl import load_workbook

    if any(name.startswith("xl/externalLinks/") for name in archive.namelist()):
        raise OfficeReadError(
            "Excel external workbook links are unsupported. Save a values-only copy before uploading."
        )
    _validate_xlsx_visibility(archive)
    text = _TextUnits()
    warnings = []
    if any(name.startswith(('xl/media/', 'xl/drawings/', 'xl/charts/')) for name in archive.namelist()):
        warnings.append(
            'Excel images, shapes and charts are not read. Review the original or upload a PDF for visual extraction.'
        )
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    cached = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    try:
        if not 1 <= len(workbook.worksheets) <= MAX_UNITS:
            raise OfficeReadError("Excel documents must contain 1–25 worksheets.")
        count = 0
        for sheet, values in zip(workbook.worksheets, cached.worksheets):
            if sheet.max_row and sheet.max_row > MAX_ROWS or sheet.max_column and sheet.max_column > 200:
                raise OfficeReadError("An Excel sheet exceeds 2,000 rows or 200 columns. Split it before uploading.")
            # Dimensions are advisory and can be intentionally false. Reset then
            # count every parsed row/cell rather than silently dropping data.
            sheet.reset_dimensions()
            values.reset_dimensions()
            for row_number, (row, cached_row) in enumerate(zip(sheet.iter_rows(), values.iter_rows()), 1):
                if row_number > MAX_ROWS or len(row) > 200:
                    raise OfficeReadError("An Excel sheet exceeds safe row or column limits.")
                cells = []
                for cell, cached_cell in zip(row, cached_row):
                    if cell.value is None:
                        continue
                    if sheet.sheet_state != 'visible':
                        raise OfficeReadError(HIDDEN_CONTENT_ERROR)
                    count += 1
                    if count > MAX_CELLS:
                        raise OfficeReadError("The workbook exceeds 10,000 populated cells. Split it before uploading.")
                    value = _display(cell.value, cell.number_format)
                    if cell.data_type == "f":
                        value = (
                            f"UNTRUSTED formula cache: {_display(cached_cell.value, cell.number_format)}"
                            if cached_cell.value is not None
                            else "UNTRUSTED formula: no cached value"
                        )
                        if not any("formula" in warning for warning in warnings):
                            warnings.append(
                                "Excel formulas were not evaluated. Cached formula results may be stale; missing results are unavailable. Review formulas in the original."
                            )
                    elif cell.data_type == "e":
                        value = f"UNTRUSTED cell error: {value}"
                    cells.append(f"{cell.coordinate}={value}")
                if cells:
                    text.add(
                        f"{sheet.title}!row {row_number}",
                        " | ".join(cells),
                        group=f"Sheet {sheet.title}",
                    )
    finally:
        workbook.close()
        cached.close()
    return text.result("xlsx", warnings)


def _xls_stream(content):
    from xlrd.compdoc import CompDoc

    compound = CompDoc(content, logfile=io.StringIO())
    for entry in compound.dirlist:
        name = entry.name.casefold()
        if any(token in name for token in ("vba", "macro", "encrypted", "encryption", "objectpool")):
            raise OfficeReadError("Encrypted, macro-enabled or embedded-object Excel documents are unsupported.")
    stream = compound.get_named_stream("Workbook") or compound.get_named_stream("Book")
    if stream is None:
        raise OfficeReadError("This legacy file is not an Excel workbook. Save Word documents as DOCX first.")
    if (
        len(stream) < 8
        or struct.unpack_from('<H', stream)[0] != 0x0809
        or struct.unpack_from('<H', stream, 4)[0] not in (0x0500, 0x0600)
    ):
        raise OfficeReadError(
            'This legacy workbook format is unsupported. Save a values-only XLSX copy before uploading.'
        )
    formulas, sheets = {}, []
    offset = 0
    while offset + 4 <= len(stream):
        kind, length = struct.unpack_from("<HH", stream, offset)
        data = stream[offset + 4 : offset + 4 + length]
        if len(data) != length:
            raise OfficeReadError("The legacy workbook contains an incomplete record.")
        if kind == 0x002F:
            raise OfficeReadError("Encrypted Excel workbooks are unsupported.")
        if kind == 0x0085 and len(data) >= 6:
            if data[5] != 0:
                raise OfficeReadError(
                    "Macro or non-worksheet legacy sheets are unsupported. Save a values-only XLSX copy."
                )
            sheets.append(struct.unpack_from("<I", data)[0])
        if kind in (0x0006, 0x0206, 0x0406) and len(data) >= 4:
            formulas[offset] = struct.unpack_from("<HH", data)
        offset += length + 4
    by_sheet = []
    for start in sheets:
        end = min((other for other in sheets if other > start), default=len(stream))
        by_sheet.append({cell for position, cell in formulas.items() if start <= position < end})
    return by_sheet


def _read_xls(content):
    import xlrd
    from openpyxl.utils.cell import get_column_letter

    formula_cells = _xls_stream(content)
    workbook = xlrd.open_workbook(file_contents=content, on_demand=True, formatting_info=True, logfile=io.StringIO())
    text = _TextUnits()
    warnings = []
    try:
        if not 1 <= workbook.nsheets <= MAX_UNITS:
            raise OfficeReadError("Excel documents must contain 1–25 worksheets.")
        count = 0
        for sheet_index in range(workbook.nsheets):
            sheet = workbook.sheet_by_index(sheet_index)
            if sheet.nrows > MAX_ROWS or sheet.ncols > 200:
                raise OfficeReadError("An Excel sheet exceeds 2,000 rows or 200 columns. Split it before uploading.")
            formulas = formula_cells[sheet_index] if sheet_index < len(formula_cells) else set()
            for row_index in range(sheet.nrows):
                cells = []
                for column_index, cell in enumerate(sheet.row(row_index)):
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    row_info = sheet.rowinfo_map.get(row_index)
                    column_info = sheet.colinfo_map.get(column_index)
                    if (
                        sheet.visibility
                        or sheet.default_row_hidden
                        or row_info is not None
                        and (row_info.hidden or row_info.height == 0)
                        or column_info is not None
                        and (column_info.hidden or column_info.width == 0)
                    ):
                        raise OfficeReadError(HIDDEN_CONTENT_ERROR)
                    count += 1
                    if count > MAX_CELLS:
                        raise OfficeReadError("The workbook exceeds 10,000 populated cells. Split it before uploading.")
                    number_format = None
                    if cell.xf_index is not None and cell.xf_index < len(workbook.xf_list):
                        format_key = workbook.xf_list[cell.xf_index].format_key
                        cell_format = workbook.format_map.get(format_key)
                        number_format = cell_format.format_str if cell_format else None
                    value = _display(cell.value, number_format)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(cell.value, workbook.datemode).isoformat()
                    if (row_index, column_index) in formulas:
                        value = f'UNTRUSTED formula cache: {value or "unavailable"}'
                        if not warnings:
                            warnings.append(
                                "Legacy Excel formulas were not evaluated. Cached results may be stale or unavailable. Review formulas in the original."
                            )
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = "UNTRUSTED cell error"
                    cells.append(f"{get_column_letter(column_index + 1)}{row_index + 1}={value}")
                if cells:
                    text.add(
                        f"{sheet.name}!row {row_index + 1}",
                        " | ".join(cells),
                        group=f"Sheet {sheet.name}",
                    )
    finally:
        workbook.release_resources()
    return text.result("xls", warnings)


def read_office(content, filename=None):
    kind = detect_intake_format(content, filename)
    if kind == "xls":
        return _read_xls(content)
    if kind not in ("docx", "xlsx"):
        raise OfficeReadError("Office text extraction requires a DOCX, XLSX or XLS document.")
    with _archive(content) as archive:
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                raw = archive.read(name).replace(b'\x00', b'').upper()
                if b"<!DOCTYPE" in raw or b"<!ENTITY" in raw:
                    raise OfficeReadError("Office XML entity declarations are not supported.")
        types = archive.read("[Content_Types].xml").lower()
        if b"macroenabled" in types or b"vbaproject" in types:
            raise OfficeReadError("Macro-enabled Office files are not supported.")
        return _read_docx(archive) if kind == "docx" else _read_xlsx(content, archive)


def main():
    # openpyxl optionally imports NumPy. Parsing cells needs no BLAS worker pool;
    # keep inherited numerical-library defaults from consuming the child budget.
    for variable in ('OPENBLAS_NUM_THREADS', 'OMP_NUM_THREADS', 'MKL_NUM_THREADS', 'NUMEXPR_NUM_THREADS'):
        os.environ[variable] = '1'
    resource.setrlimit(resource.RLIMIT_CPU, (15, 15))
    if sys.platform.startswith("linux"):
        resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
    try:
        content = sys.stdin.buffer.read(MAX_BYTES + 1)
        result = read_office(content, sys.argv[1] if len(sys.argv) > 1 else None)
        print(json.dumps(result))
    except Exception as exc:
        message = (
            str(exc)
            if isinstance(exc, OfficeReadError)
            else "The Office document is damaged or could not be safely read."
        )
        print(json.dumps({"error": message}))
        raise SystemExit(1) from None


if __name__ == "__main__":
    main()
