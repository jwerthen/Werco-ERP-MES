"""Native Office extraction preserves evidence and refuses unsafe or partial input."""

import io
import struct
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook

from app.services.hank_office_reader import OfficeReadError, detect_intake_format, read_office


def word_document():
    document = Document()
    document.add_paragraph('Purchase Order PO-WORD-123')
    document.add_paragraph('Supplier: Reliable Metals')
    table = document.add_table(rows=1, cols=4)
    for cell, value in zip(table.rows[0].cells, ['Part', 'Quantity', 'Unit', 'Unit price']):
        cell.text = value
    for cell, value in zip(table.add_row().cells, ['MAT-001', '12', 'EA', '4.25']):
        cell.text = value
    document.add_paragraph('Delivery requested 2026-10-01')
    result = io.BytesIO()
    document.save(result)
    return result.getvalue()


def excel_document(*, formula=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Material PO'
    sheet.append(['Purchase Order', 'PO-EXCEL-123'])
    sheet.append(['Part', 'Quantity', 'Unit', 'Unit price'])
    sheet.append(['MAT-001', '=6*2' if formula else 12, 'EA', 4.25])
    result = io.BytesIO()
    workbook.save(result)
    return result.getvalue()


def legacy_excel_document(*, formula=False, macro=False, encrypted=False, hidden=None):
    """Minimal inert OLE/BIFF8 fixture, with no optional writer dependency."""

    def record(kind, body=b''):
        return struct.pack('<HH', kind, len(body)) + body

    def bof(kind):
        return record(0x0809, struct.pack('<HHHHII', 0x0600, kind, 3515, 1996, 0, 6))

    strings = ['PO-XLS-123', 'MAT-001', 'EA']
    sst = record(
        0x00FC, struct.pack('<II', 3, 3) + b''.join(struct.pack('<HB', len(s), 0) + s.encode() for s in strings)
    )
    boundsheet = lambda offset: record(
        0x0085, struct.pack('<IBBBB', offset, int(hidden == 'sheet'), int(macro), 2, 0) + b'PO'
    )
    prefix = bof(5) + record(0x0042, struct.pack('<H', 1200))
    if encrypted:
        prefix += record(0x002F, b'\0' * 6)
    globals_size = len(prefix + boundsheet(0) + sst + record(0x000A))
    stream = prefix + boundsheet(globals_size) + sst + record(0x000A) + bof(16)
    stream += record(0x0200, struct.pack('<IIHHH', 0, 2, 0, 3, 0))
    if hidden == 'row':
        stream += record(0x0208, struct.pack('<HHHHHHI', 1, 0, 3, 255, 0, 0, 0x20))
    if hidden == 'column':
        stream += record(0x007D, struct.pack('<HHHHHH', 0, 0, 2048, 0, 1, 0))
    if hidden == 'default_row':
        stream += record(0x0225, struct.pack('<HH', 2, 255))
    for row, column, index in [(0, 0, 0), (1, 0, 1), (1, 2, 2)]:
        stream += record(0x00FD, struct.pack('<HHHI', row, column, 0, index))
    if formula:
        stream += record(0x0006, struct.pack('<HHHdH4sH', 1, 1, 0, 12, 0, b'\0' * 4, 3) + b'\x1e\x0c\x00')
    else:
        stream += record(0x0203, struct.pack('<HHHd', 1, 1, 0, 12))
    stream = (stream + record(0x000A)).ljust(4096, b'\0')
    header = bytearray(512)
    header[:8] = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'
    struct.pack_into('<HHHH', header, 24, 0x3E, 3, 0xFFFE, 9)
    struct.pack_into('<H', header, 32, 6)
    struct.pack_into('<IIIIIIIII', header, 40, 0, 1, 0, 0, 4096, 0xFFFFFFFE, 0, 0xFFFFFFFE, 0)
    struct.pack_into('<109I', header, 76, 9, *([0xFFFFFFFF] * 108))

    def directory(name, kind, child, sector, size):
        entry = bytearray(128)
        encoded = (name + '\0').encode('utf-16le')
        entry[: len(encoded)] = encoded
        struct.pack_into('<HBBiii', entry, 64, len(encoded), kind, 1, -1, -1, child)
        struct.pack_into('<IQ', entry, 116, sector, size)
        return entry

    entries = directory('Root Entry', 5, 1, 0xFFFFFFFE, 0) + directory('Workbook', 2, -1, 1, 4096) + bytes(256)
    fat = struct.pack('<128I', 0xFFFFFFFE, *range(2, 9), 0xFFFFFFFE, 0xFFFFFFFD, *([0xFFFFFFFF] * 118))
    return bytes(header) + entries + stream + fat


def rewrite_archive(content, updates):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as source, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
        for name in source.namelist():
            if name not in updates:
                target.writestr(name, source.read(name))
        for name, value in updates.items():
            target.writestr(name, value)
    return output.getvalue()


def test_word_preserves_body_paragraph_table_order_and_locators():
    parsed = read_office(word_document(), 'purchase-order.docx')
    text = '\n'.join(parsed['units'])
    assert parsed['format'] == 'docx'
    assert text.index('PO-WORD-123') < text.index('MAT-001') < text.index('Delivery requested')
    assert '[Document table 1 row 2] MAT-001 | 12 | EA | 4.25' in text
    assert parsed['labels'] == ['Document — section 1']


@pytest.mark.parametrize(
    'revision',
    [
        '<w:trPr><w:del/></w:trPr>',
        '<w:tcPr><w:cellDel/></w:tcPr>',
        '<w:pPr><w:rPr><w:del/></w:rPr></w:pPr>',
        '<w:moveFromRangeStart w:id="1"/>',
    ],
)
def test_word_structural_tracked_revisions_are_rejected_instead_of_importing_deleted_lines(revision):
    content = word_document()
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        xml = archive.read('word/document.xml').replace(b'<w:tr>', ('<w:tr>' + revision).encode(), 1)
    content = rewrite_archive(content, {'word/document.xml': xml})
    with pytest.raises(OfficeReadError, match='tracked changes'):
        read_office(content, 'PO.docx')


@pytest.mark.parametrize('factory,extension', [(excel_document, 'xlsx'), (legacy_excel_document, 'xls')])
def test_excel_reads_actual_cells_and_preserves_sheet_evidence(factory, extension):
    parsed = read_office(factory(), f'purchase-order.{extension}')
    assert parsed['format'] == extension
    text = '\n'.join(parsed['units'])
    assert 'A' in text and '=MAT-001' in text and '=12' in text and '=EA' in text
    assert '!row ' in text


@pytest.mark.parametrize('factory,extension', [(excel_document, 'xlsx'), (legacy_excel_document, 'xls')])
def test_excel_formulas_are_not_evaluated_and_caches_are_untrusted(factory, extension):
    parsed = read_office(factory(formula=True), f'purchase-order.{extension}')
    assert 'UNTRUSTED formula' in '\n'.join(parsed['units'])
    assert any('not evaluated' in warning for warning in parsed['warnings'])


@pytest.mark.parametrize(
    'content,filename',
    [
        (b'not office', 'po.docx'),
        (b'%PDF-pretend', 'po.xlsx'),
        (word_document(), 'po.xlsx'),
        (excel_document(), 'po.docx'),
    ],
)
def test_disguised_files_are_rejected(content, filename):
    with pytest.raises(OfficeReadError):
        detect_intake_format(content, filename)


@pytest.mark.parametrize(
    'name,payload',
    [
        ('word/vbaProject.bin', b'macro'),
        ('word/embeddings/oleObject1.bin', b'object'),
        ('word/document.xml', b'<!DOCTYPE x [<!ENTITY x SYSTEM "file:///etc/passwd">]><x>&x;</x>'),
        ('huge.txt', b'a' * (2 * 1024 * 1024)),
    ],
)
def test_unsafe_office_packages_are_rejected(name, payload):
    with pytest.raises(OfficeReadError):
        read_office(rewrite_archive(word_document(), {name: payload}), 'po.docx')


def test_external_excel_links_and_legacy_macros_encryption_are_rejected():
    with pytest.raises(OfficeReadError, match='external'):
        read_office(rewrite_archive(excel_document(), {'xl/externalLinks/externalLink1.xml': b'<links/>'}), 'po.xlsx')
    for option in ('macro', 'encrypted'):
        with pytest.raises(OfficeReadError):
            read_office(legacy_excel_document(**{option: True}), 'po.xls')


def test_excessive_word_text_is_rejected_without_truncation():
    document = Document()
    document.add_paragraph('A' * 17000)
    stream = io.BytesIO()
    document.save(stream)
    with pytest.raises(OfficeReadError, match='too much text'):
        read_office(stream.getvalue(), 'po.docx')


def test_excel_false_dimensions_do_not_hide_material_rows():
    workbook = Workbook()
    workbook.active['A1'] = 'PO-123'
    hidden = workbook.create_sheet('Hidden material')
    hidden['A7'] = 'MAT-SECRET'
    stream = io.BytesIO()
    workbook.save(stream)
    content = stream.getvalue()
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        hidden_xml = archive.read('xl/worksheets/sheet2.xml').replace(b'A7:A7', b'A1:A1')
    parsed = read_office(rewrite_archive(content, {'xl/worksheets/sheet2.xml': hidden_xml}), 'po.xlsx')
    assert any('MAT-SECRET' in unit for unit in parsed['units'])


@pytest.mark.parametrize('hidden', ['sheet', 'row', 'column', 'default_row'])
def test_populated_hidden_legacy_excel_content_is_rejected(hidden):
    with pytest.raises(OfficeReadError, match='populated hidden'):
        read_office(legacy_excel_document(hidden=hidden), 'PO.xls')


@pytest.mark.parametrize('hidden', ['sheet', 'row', 'column', 'zero_height', 'default_hidden'])
def test_populated_hidden_excel_content_is_rejected(hidden):
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(excel_document()))
    sheet = workbook.active
    if hidden == 'sheet':
        workbook.create_sheet('Visible')
        sheet.sheet_state = 'hidden'
    elif hidden == 'row':
        sheet.row_dimensions[3].hidden = True
    elif hidden == 'column':
        sheet.column_dimensions['A'].hidden = True
    elif hidden == 'zero_height':
        sheet.row_dimensions[3].height = 0
    else:
        sheet.sheet_format.zeroHeight = True
    stream = io.BytesIO()
    workbook.save(stream)
    content = stream.getvalue()
    if hidden == 'zero_height':
        # openpyxl omits ht=0 while authoring; explicitly zero-height OOXML is
        # meaningful input from another producer and must not hide a PO line.
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            xml = archive.read('xl/worksheets/sheet1.xml').replace(b'<row r="3"', b'<row r="3" ht="0"')
        content = rewrite_archive(content, {'xl/worksheets/sheet1.xml': xml})
    with pytest.raises(OfficeReadError, match='populated hidden'):
        read_office(content, 'PO.xlsx')


def test_empty_hidden_excel_formatting_is_allowed():
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(excel_document()))
    workbook.active.row_dimensions[20].hidden = True
    workbook.active.column_dimensions['Z'].hidden = True
    workbook.create_sheet('Empty hidden').sheet_state = 'hidden'
    stream = io.BytesIO()
    workbook.save(stream)
    assert 'MAT-001' in '\n'.join(read_office(stream.getvalue(), 'PO.xlsx')['units'])


def test_excel_limit_rejects_relevant_rows_beyond_bound():
    workbook = Workbook()
    workbook.active['A2001'] = 'MAT-TOO-FAR'
    stream = io.BytesIO()
    workbook.save(stream)
    with pytest.raises(OfficeReadError, match='2,000 rows'):
        read_office(stream.getvalue(), 'po.xlsx')
